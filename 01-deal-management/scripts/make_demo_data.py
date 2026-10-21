"""
案件管理アプリのデモデータ（kintone インポート用 Excel）を作るスクリプト。

次回アクション期限は「実行した日」を基準に作るため、
インポート直前に実行し直すと、期限切れ・期限間近の色分けがいつでも再現できます。

使い方: python3 make_demo_data.py [自社担当者のログイン名]
"""
import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

SALES_REP = sys.argv[1] if len(sys.argv) > 1 else 'ログイン名に置き換え'
TODAY = date.today()
OUT = Path(__file__).resolve().parent.parent / 'demo-data' / 'deal_demo_data.xlsx'


def d(offset):
    """今日から offset 日後の日付を 'YYYY-MM-DD' で返す（None なら空欄）"""
    return '' if offset is None else (TODAY + timedelta(days=offset)).isoformat()


# 1行目の見出しはフィールド名と同じにしておくと、インポート時に自動で対応付けされる
HEADERS = [
    '案件名', '顧客名', '顧客担当者', '電話番号', 'メールアドレス', '自社担当者',
    'ステータス', '確度', '見込み金額', '次回アクション期限', '次回アクション内容',
    '提案日', '見積金額', '見積提出日', '受注日', '受注金額', '失注理由', '失注理由の詳細', '備考',
]

# (案件名, 顧客名, 顧客担当者, ステータス, 確度, 見込み金額, 期限(日後), 次回アクション内容,
#  提案日(日後), 見積金額, 見積提出日(日後), 受注日(日後), 受注金額, 失注理由, 失注理由の詳細, 備考)
DEALS = [
    # --- 期限切れ（赤）---
    ('勤怠管理システム導入', '株式会社サンプル商事', '山田 太郎', '見積提出', 'A', 1800000, -5, '見積内容の回答確認の電話', -30, 1800000, -14, None, None, '', '', '決裁者は部長'),
    ('営業日報のクラウド化', '有限会社みどり運輸', '佐藤 花子', '提案中', 'B', 600000, -2, '提案書の修正版を送付', -10, None, None, None, None, '', '', ''),
    ('Webサイトリニューアル', '株式会社ひかり不動産', '鈴木 一郎', 'リード', 'C', 1200000, -8, '初回ヒアリングの日程調整', None, None, None, None, None, '', '', '展示会で名刺交換'),
    ('在庫管理アプリ構築', '株式会社北斗食品', '高橋 美咲', '見積提出', 'B', 2400000, -1, '相見積もりの状況を確認', -25, 2350000, -7, None, None, '', '', ''),
    # --- 期限間近（黄）---
    ('問い合わせ管理の効率化', '株式会社あおば製作所', '田中 健', '提案中', 'A', 900000, 0, 'デモ環境の案内', -6, None, None, None, None, '', '', '本日中に連絡'),
    ('契約書管理の電子化', '税理士法人さくら会計', '伊藤 直子', '見積提出', 'A', 750000, 1, '契約条件のすり合わせ', -20, 720000, -5, None, None, '', '', ''),
    ('採用管理アプリ導入', '株式会社ミライ人材', '渡辺 翔', 'リード', 'B', 500000, 2, '課題ヒアリング', None, None, None, None, None, '', '', ''),
    ('設備点検記録のデジタル化', '株式会社東和ビルメンテ', '中村 優子', '提案中', 'B', 1500000, 3, '現場担当者との打ち合わせ', -4, None, None, None, None, '', '', ''),
    # --- 期限に余裕あり ---
    ('顧客台帳の一元化', '株式会社やまと工務店', '小林 誠', 'リード', 'C', 400000, 10, '資料送付後のフォロー', None, None, None, None, None, '', '', ''),
    ('経費精算フローの改善', '株式会社こだま印刷', '加藤 由美', '提案中', 'B', 650000, 7, '提案内容の社内共有待ち', -3, None, None, None, None, '', '', ''),
    ('研修申込管理アプリ', '一般社団法人みらい教育協会', '吉田 拓也', '見積提出', 'A', 1100000, 14, '理事会での承認待ち', -21, 1050000, -9, None, None, '', '', '理事会は月末'),
    ('点検報告書の自動作成', '株式会社しらかわ電設', '山本 彩', 'リード', 'B', 800000, 21, 'オンライン面談', None, None, None, None, None, '', '', ''),
    ('クレーム対応履歴の管理', '株式会社ふじ観光', '松本 大輔', '提案中', 'C', 550000, 12, 'ヒアリング内容をもとに提案書作成', -2, None, None, None, None, '', '', ''),
    ('会員管理システム刷新', 'NPO法人つなぐ', '井上 麻衣', 'リード', 'C', 300000, 30, '予算確定後に再連絡', None, None, None, None, None, '', '', '来期予算で検討'),
    # --- 受注 ---
    ('案件管理アプリ構築', '株式会社ことぶき商会', '木村 亮', '受注', 'A', 1600000, None, '', -60, 1600000, -45, -30, 1600000, '', '', ''),
    ('日報アプリ導入支援', '株式会社はやて物流', '林 さやか', '受注', 'A', 450000, None, '', -40, 480000, -35, -20, 450000, '', '', '値引き対応あり'),
    ('見積・請求の連携', '株式会社オリオン機器', '清水 修', '受注', 'A', 2000000, None, '', -90, 2100000, -70, -50, 2000000, '', '', ''),
    ('アンケート集計の自動化', '株式会社ひまわり保育', '山口 恵', '受注', 'B', 350000, None, '', -15, 350000, -10, -3, 350000, '', '', ''),
    ('工事進捗の見える化', '株式会社いずみ建設', '森 浩二', '受注', 'A', 2800000, None, '', -100, 2800000, -80, -60, 2750000, '', '', ''),
    # --- 失注 ---
    ('予約管理システム', '株式会社あさひクリニック', '池田 真由美', '失注', 'B', 900000, None, '', -50, 950000, -40, None, None, '価格', '', ''),
    ('販売管理の刷新', '株式会社たちばな商店', '橋本 剛', '失注', 'C', 1300000, None, '', -70, 1400000, -55, None, None, '競合他社', '', ''),
    ('社内ポータル構築', '株式会社ゆうき精機', '石川 理恵', '失注', 'B', 700000, None, '', -45, None, None, None, None, '時期・予算', '', ''),
    ('入退室管理の連携', '株式会社かなえ警備', '前田 隆', '失注', 'C', 1000000, None, '', -35, 1000000, -28, None, None, 'その他', '先方の担当者が異動し、プロジェクトが凍結', ''),
]

wb = Workbook()
ws = wb.active
ws.title = 'インポート用'
ws.append(HEADERS)

for i, (name, customer, contact, status, prob, expected, due, action,
        proposal, quote_amt, quote_date, order_date, order_amt,
        lost, lost_detail, memo) in enumerate(DEALS, start=1):
    ws.append([
        name, customer, contact, f'03-0000-{i:04d}', f'demo{i:02d}@example.com', SALES_REP,
        status, prob, expected, d(due), action,
        d(proposal), quote_amt, d(quote_date), d(order_date), order_amt, lost, lost_detail, memo,
    ])

# 見た目の調整（インポートには影響しない）
header_fill = PatternFill('solid', fgColor='DDEBF7')
for cell in ws[1]:
    cell.font = Font(name='Arial', bold=True)
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center')
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.font = Font(name='Arial')
        # 日付・電話番号は文字列として扱う（Excel に日付や数値へ変換させない）
        if isinstance(cell.value, str):
            cell.number_format = '@'
        elif isinstance(cell.value, int):
            cell.number_format = '#,##0'
widths = [24, 26, 14, 15, 24, 28, 10, 6, 12, 16, 30, 12, 12, 12, 12, 12, 12, 34, 18]
for col, w in zip('ABCDEFGHIJKLMNOPQRS', widths):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A2'

# 2枚目：使い方のメモ（インポートは1枚目のシートだけ読み込まれる）
note = wb.create_sheet('使い方')
for line in [
    ['このファイルについて'],
    [f'作成日：{TODAY.isoformat()}（次回アクション期限はこの日を基準にしています）'],
    ['インポート：アプリの一覧画面 →「…」→「ファイルから読み込む」で、このファイルを選択'],
    ['1枚目「インポート用」だけを読み込みます。1行目は見出しなので「先頭行は見出しとして扱う」を選択'],
    ['「自社担当者」列は kintone のログイン名です。違う場合は一括置換してください'],
    ['期限の色分けを再現したいときは scripts/make_demo_data.py を実行し直してから読み込んでください'],
]:
    note.append(line)
note.column_dimensions['A'].width = 100
for row in note.iter_rows():
    for cell in row:
        cell.font = Font(name='Arial', bold=(cell.row == 1))

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f'{len(DEALS)}件を書き出しました: {OUT}')
